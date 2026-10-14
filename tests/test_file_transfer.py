"""Moving whole files: file resources, the file manager panel and its upload
tools, download links, and what read_file/write_file tell agents that can't
use them."""

import base64
import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

import fulcra_mcp.file_transfer as ft
from fulcra_mcp.settings import settings

FILE = {"id": "f-1", "name": "q3.xlsx", "path": "/reports", "size": 11}


@pytest.fixture(autouse=True)
def state_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "state_path", tmp_path / "state")
    return tmp_path / "state"


def _download_returns(fake_fulcra, data: bytes):
    fake_fulcra.download_file.return_value = io.BytesIO(data)


def _xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Totals"
    sheet.append(["region", "sales"])
    sheet.append(["west", 12])
    other = workbook.create_sheet("Notes")
    other.append(['say "hi", ok'])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


# --- resources -------------------------------------------------------------------


async def test_a_file_is_served_as_a_blob_resource(client, fake_fulcra):
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, size=4)]
    _download_returns(fake_fulcra, b"\x00\x01\xff\xfe")

    [content] = await client.read_resource("fulcra://files/reports/q3%20final.xlsx")

    fake_fulcra.resolve_filepath.assert_called_once_with(
        "/reports/q3 final.xlsx", fulcra_userid=None
    )
    assert base64.b64decode(content.blob) == b"\x00\x01\xff\xfe"
    assert content.mimeType == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


async def test_a_resource_path_is_decoded_once(client, fake_fulcra):
    # A file literally named "literal%20name.txt" is encoded as %2520.
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, name="literal%20name.txt", size=2)]
    _download_returns(fake_fulcra, b"hi")

    await client.read_resource(ft.file_resource_uri("/docs/literal%20name.txt"))

    fake_fulcra.resolve_filepath.assert_called_once_with(
        "/docs/literal%20name.txt", fulcra_userid=None
    )


async def test_a_shared_file_resource_reads_as_that_user(client, fake_fulcra):
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, name="a.txt", size=2)]
    _download_returns(fake_fulcra, b"hi")

    [content] = await client.read_resource("fulcra://users/u-2/files/notes/a.txt")

    fake_fulcra.resolve_filepath.assert_called_once_with("/notes/a.txt", fulcra_userid="u-2")
    assert content.text == "hi"


async def test_a_file_too_large_for_one_message_points_to_a_link(client, fake_fulcra):
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, size=ft.MAX_RESOURCE_BYTES + 1)]

    with pytest.raises(Exception, match="as_link=true"):
        await client.read_resource("fulcra://files/reports/q3.xlsx")
    fake_fulcra.download_file.assert_not_called()


# --- the panel -------------------------------------------------------------------


async def test_the_panel_is_an_mcp_app_resource(client, fake_fulcra):
    [content] = await client.read_resource(ft.FILE_MANAGER_URI)

    assert content.mimeType == "text/html;profile=mcp-app"
    # Rubik from Google Fonts, allowed by the panel's CSP
    [resource] = [r for r in await client.list_resources() if str(r.uri) == ft.FILE_MANAGER_URI]
    assert resource.meta["ui"]["csp"]["resourceDomains"] == [
        "https://fonts.googleapis.com",
        "https://fonts.gstatic.com",
    ]
    for method in ("ui/initialize", "resources/read", "ui/download-file", "upload_file_chunk"):
        assert method in content.text


async def test_tools_declare_the_panel_and_app_only_visibility(client, fake_fulcra):
    tools = {t.name: t for t in await client.list_tools()}

    assert tools["open_file_manager"].meta["ui"] == {"resourceUri": ft.FILE_MANAGER_URI}
    for name in ("start_file_upload", "upload_file_chunk"):
        # Listed, so hosts can let the panel call them, but hidden from the model.
        assert tools[name].meta["ui"]["visibility"] == ["app"]
    assert "get_file_download_link" not in tools  # folded into read_file(as_link)


async def test_the_files_prompt_is_listed_with_an_optional_path(client, fake_fulcra):
    [prompt] = [p for p in await client.list_prompts() if p.name == "files"]

    assert prompt.title == "Fulcra files"
    assert [(a.name, a.required) for a in prompt.arguments] == [("path", False)]


@pytest.mark.parametrize("args, folder", [({}, "/"), ({"path": "reports/2026"}, "/reports/2026")])
async def test_the_files_prompt_asks_for_the_file_manager(client, fake_fulcra, args, folder):
    result = await client.get_prompt("files", args)

    [message] = result.messages
    assert message.role == "user"
    assert message.content.text == (
        f"Open my Fulcra file manager at {folder} with open_file_manager, "
        "so I can upload or download files."
    )


async def test_open_file_manager_helps_clients_without_the_panel(call, fake_fulcra):
    fake_fulcra.list_files.return_value = {"folders": ["/reports/2026"], "files": [FILE]}

    text = await call("open_file_manager", {"path": "reports"})

    fake_fulcra.list_files.assert_called_once_with("/reports")
    assert "write_file" in text and "as_link=true" in text
    assert '"q3.xlsx"' in text


# --- uploads ---------------------------------------------------------------------


@pytest.fixture
def put_chunks(monkeypatch):
    """Stand in for the storage service: commits every chunk it's given."""
    sent = []

    def fake_put(url, data, start, total):
        sent.append((url, data, start, total))
        return start + len(data)

    monkeypatch.setattr(ft, "_put_chunk", fake_put)
    return sent


def _start(fake_fulcra, size):
    fake_fulcra.fulcra_api.return_value = json.dumps(
        {"url": "https://storage.example/upload?id=1", "file": dict(FILE, size=size)}
    ).encode()


async def _structured(client, name, args):
    result = await client.call_tool(name, args)
    return result.structured_content


async def test_an_upload_goes_through_in_chunks(client, fake_fulcra, put_chunks):
    size = ft.UPLOAD_GRANULE + 5
    _start(fake_fulcra, size)
    data = bytes(range(256)) * (size // 256) + bytes(size % 256)

    started = await _structured(
        client, "start_file_upload", {"path": "/reports/q3.xlsx", "size": size}
    )
    fake_fulcra.fulcra_api.assert_called_once_with(
        "/input/v1/file_upload",
        method="POST",
        data={
            "content_length": size,
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "name": "q3.xlsx",
            "path": "/reports",
        },
    )
    first = await _structured(
        client,
        "upload_file_chunk",
        {
            "upload_id": started["upload_id"],
            "offset": 0,
            "data": base64.b64encode(data[: ft.UPLOAD_GRANULE]).decode(),
        },
    )
    assert (first["received"], first["done"]) == (ft.UPLOAD_GRANULE, False)

    last = await _structured(
        client,
        "upload_file_chunk",
        {
            "upload_id": started["upload_id"],
            "offset": ft.UPLOAD_GRANULE,
            "data": base64.b64encode(data[ft.UPLOAD_GRANULE :]).decode(),
        },
    )
    assert last["done"] and last["path"] == "/reports/q3.xlsx"
    assert last["resource_uri"] == "fulcra://files/reports/q3.xlsx"
    assert b"".join(chunk for _, chunk, _, _ in put_chunks) == data
    assert [(start, total) for _, _, start, total in put_chunks] == [
        (0, size),
        (ft.UPLOAD_GRANULE, size),
    ]


@pytest.mark.parametrize(
    "offset, chunk, message",
    [
        (0, b"x" * 1000, "multiple of 262144"),
        (5, b"x" * ft.UPLOAD_GRANULE, "Expected the chunk starting at byte 0"),
        (0, b"x" * (2 * ft.UPLOAD_GRANULE + 1), "past the file's size"),
    ],
)
async def test_bad_chunks_are_refused(client, fake_fulcra, put_chunks, offset, chunk, message):
    _start(fake_fulcra, 2 * ft.UPLOAD_GRANULE)
    started = await _structured(client, "start_file_upload", {"path": "/a.bin", "size": 2 * ft.UPLOAD_GRANULE})

    result = await client.call_tool(
        "upload_file_chunk",
        {"upload_id": started["upload_id"], "offset": offset, "data": base64.b64encode(chunk).decode()},
        raise_on_error=False,
    )

    assert result.is_error and message in result.content[0].text
    assert put_chunks == []


async def test_a_resent_chunk_is_harmless(client, fake_fulcra, put_chunks):
    _start(fake_fulcra, 2 * ft.UPLOAD_GRANULE)
    started = await _structured(client, "start_file_upload", {"path": "/a.bin", "size": 2 * ft.UPLOAD_GRANULE})
    chunk = {"upload_id": started["upload_id"], "offset": 0, "data": base64.b64encode(b"x" * ft.UPLOAD_GRANULE).decode()}

    await client.call_tool("upload_file_chunk", chunk)
    again = await _structured(client, "upload_file_chunk", chunk)

    assert again["received"] == ft.UPLOAD_GRANULE and len(put_chunks) == 1


async def test_resending_the_last_chunk_returns_the_same_result(client, fake_fulcra, put_chunks):
    _start(fake_fulcra, 3)
    started = await _structured(client, "start_file_upload", {"path": "/a.bin", "size": 3})
    last = {"upload_id": started["upload_id"], "offset": 0, "data": base64.b64encode(b"abc").decode()}

    first = await _structured(client, "upload_file_chunk", last)
    again = await _structured(client, "upload_file_chunk", last)

    assert first["done"] and again == first
    assert len(put_chunks) == 1  # stored once


async def test_resending_an_empty_files_only_chunk_returns_the_same_result(client, fake_fulcra, put_chunks):
    _start(fake_fulcra, 0)
    started = await _structured(client, "start_file_upload", {"path": "/empty.txt", "size": 0})
    only = {"upload_id": started["upload_id"], "offset": 0, "data": ""}

    first = await _structured(client, "upload_file_chunk", only)
    again = await _structured(client, "upload_file_chunk", only)

    assert first["done"] and again == first
    assert len(put_chunks) == 1


async def test_an_empty_file_uploads(client, fake_fulcra, put_chunks):
    _start(fake_fulcra, 0)
    started = await _structured(client, "start_file_upload", {"path": "/empty.txt", "size": 0})

    done = await _structured(client, "upload_file_chunk", {"upload_id": started["upload_id"], "offset": 0, "data": ""})

    assert done["done"] and put_chunks[0][1:] == (b"", 0, 0)


async def test_someone_elses_upload_is_unknown(client, fake_fulcra, put_chunks, monkeypatch):
    _start(fake_fulcra, 1)
    started = await _structured(client, "start_file_upload", {"path": "/a.bin", "size": 1})
    monkeypatch.setattr(ft, "current_grant_id", lambda: "another-login")

    result = await client.call_tool(
        "upload_file_chunk",
        {"upload_id": started["upload_id"], "offset": 0, "data": "eA=="},
        raise_on_error=False,
    )

    assert result.is_error and "No upload" in result.content[0].text


def test_put_chunk_reads_partial_progress(monkeypatch):
    import urllib.error

    def resume_incomplete(request):
        assert request.get_header("Content-range") == "bytes 0-3/10"
        raise urllib.error.HTTPError(request.full_url, 308, "Resume", {"Range": "bytes=0-3"}, None)

    monkeypatch.setattr(ft.urllib.request, "urlopen", resume_incomplete)
    assert ft._put_chunk("https://storage.example/u", b"abcd", 0, 10) == 4


# --- download links ----------------------------------------------------------------


async def test_a_local_server_saves_the_file_to_downloads(call, fake_fulcra, tmp_path, monkeypatch):
    monkeypatch.setattr(ft.Path, "home", lambda: tmp_path)
    (tmp_path / "Downloads").mkdir()
    (tmp_path / "Downloads" / "q3.xlsx").write_bytes(b"old")
    fake_fulcra.resolve_filepath.return_value = [FILE]
    _download_returns(fake_fulcra, b"spreadsheet")

    text = await call("read_file", {"path": "/reports/q3.xlsx", "as_link": True})

    saved = tmp_path / "Downloads" / "q3 (1).xlsx"
    assert saved.read_bytes() == b"spreadsheet"
    assert str(saved) in text


@pytest.fixture
def hosted(monkeypatch):
    monkeypatch.setattr(settings, "fulcra_environment", "hosted")
    monkeypatch.setattr(settings, "oidc_server_url", "https://mcp.example")
    monkeypatch.setattr(ft, "current_grant_id", lambda: "grant-1")


async def test_a_hosted_server_makes_a_short_lived_link(call, fake_fulcra, hosted):
    fake_fulcra.resolve_filepath.return_value = [FILE]

    link = json.loads(await call("read_file", {"path": "/reports/q3.xlsx", "as_link": True}))

    assert link["url"].startswith("https://mcp.example/files/download/")
    token = link["url"].rsplit("/", 1)[1]
    target = ft.download_link_target(token)
    assert target["grant_id"] == "grant-1" and target["file_id"] == "f-1"
    # only a hash of the token is stored
    assert token not in "".join(p.name for p in settings.state_path.rglob("*"))


def test_the_download_route_streams_the_file(fake_fulcra, monkeypatch):
    import fulcra_mcp.main as main

    ft._save_record(
        "download_links",
        "tok",
        {"grant_id": "g", "file_id": "f-1", "fulcra_userid": None, "name": "Q3 résumé.xlsx", "expires_at": time.time() + 60},
    )
    monkeypatch.setattr(main.oauth_provider, "reload_grant", lambda grant_id: object())
    monkeypatch.setattr(main, "fulcra_for_grant", lambda grant_id, creds: fake_fulcra)
    upstream = io.BytesIO(b"spreadsheet")
    upstream.headers = {"Content-Length": "11"}
    fake_fulcra.download_file.return_value = upstream

    response = TestClient(main.app).get("/files/download/tok")

    assert response.status_code == 200 and response.content == b"spreadsheet"
    assert "attachment" in response.headers["content-disposition"]
    assert "filename*=UTF-8''Q3%20r%C3%A9sum%C3%A9.xlsx" in response.headers["content-disposition"]
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")


def test_an_expired_link_says_so(fake_fulcra):
    import fulcra_mcp.main as main

    ft._save_record("download_links", "old", {"grant_id": "g", "file_id": "f", "name": "a", "expires_at": time.time() - 1})

    response = TestClient(main.app).get("/files/download/old")

    assert response.status_code == 410 and "expired" in response.text
    assert not list((settings.state_path / "download_links").glob("*.json"))


# --- read_file and write_file --------------------------------------------------------


async def test_read_file_returns_a_spreadsheet_as_csv(call, fake_fulcra):
    fake_fulcra.resolve_filepath.return_value = [FILE]
    _download_returns(fake_fulcra, _xlsx())

    text = await call("read_file", {"path": "/reports/q3.xlsx"})

    assert "(as CSV" in text
    assert "# Sheet: Totals\nregion,sales\nwest,12" in text
    assert '"say ""hi"", ok"' in text


class LimitedResponse(io.BytesIO):
    """A download that records how much was asked of it."""

    def __init__(self, data):
        super().__init__(data)
        self.asked = []

    def read(self, size=-1):
        self.asked.append(size)
        return super().read(size)


def _sparse_xlsx(far_cell: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "first"
    sheet[far_cell] = "last"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


async def test_a_workbook_spanning_the_whole_grid_converts_within_bounds(call, fake_fulcra):
    data = _sparse_xlsx("XFD1048576")  # Excel's last cell
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, size=len(data))]
    _download_returns(fake_fulcra, data)

    started = time.monotonic()
    text = await call("read_file", {"path": "/reports/q3.xlsx", "max_bytes": 5000})

    assert time.monotonic() - started < 10
    assert len(text) < 5500
    assert text.count("\n") < 50  # empty rows and trailing cells are left out
    assert "only the first 200 columns" in text
    assert "stopped after" in text


def test_conversion_stops_at_the_text_limit():
    workbook = Workbook()
    for i in range(5000):
        workbook.active.append([f"row {i}", i])
    out = io.BytesIO()
    workbook.save(out)

    text = ft.spreadsheet_as_csv(out.getvalue(), 1000)

    assert "row 0,0" in text and "row 4999" not in text
    assert "stopped after the text limit" in text
    assert len(text) < 1100


async def test_a_large_spreadsheet_is_refused_before_downloading(call, fake_fulcra):
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, size=ft.MAX_SPREADSHEET_FILE_BYTES + 1)]

    text = await call("read_file", {"path": "/reports/q3.xlsx"})

    assert "too large to convert" in text and "as_link=true" in text
    fake_fulcra.download_file.assert_not_called()


async def test_conversion_runs_off_the_event_loop(call, fake_fulcra, monkeypatch):
    import fulcra_mcp.tools as tools_module

    fake_fulcra.resolve_filepath.return_value = [FILE]
    _download_returns(fake_fulcra, _xlsx())
    ran = []
    real = tools_module.asyncio.to_thread

    async def recording(func, *args):
        ran.append(func.__name__)
        return await real(func, *args)

    monkeypatch.setattr(tools_module.asyncio, "to_thread", recording)
    await call("read_file", {"path": "/reports/q3.xlsx"})

    assert ran == ["spreadsheet_as_csv"]


async def test_text_is_read_only_as_far_as_max_bytes(call, fake_fulcra):
    response = LimitedResponse("é".encode() * 50_000)  # 100,000 bytes
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, name="a.txt", size=100_000)]
    fake_fulcra.download_file.return_value = response

    text = await call("read_file", {"path": "/a.txt", "max_bytes": 1001})

    assert response.asked == [1005]
    assert "truncated to the first 1001 bytes of 100000 bytes" in text
    # the character cut in half at the limit is dropped, not garbled
    assert text.endswith("é" * 500)


async def test_max_bytes_is_capped(call, fake_fulcra):
    response = LimitedResponse(b"x" * 10)
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, name="a.txt", size=10)]
    fake_fulcra.download_file.return_value = response

    await call("read_file", {"path": "/a.txt", "max_bytes": 10**12})

    assert response.asked == [ft.MAX_RESOURCE_BYTES + 4]


async def test_large_binary_with_include_binary_is_refused_after_a_bounded_read(call, fake_fulcra):
    response = LimitedResponse(b"\xff\xfe" * 500_000)
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, name="a.bin", size=1_000_000)]
    fake_fulcra.download_file.return_value = response

    text = await call("read_file", {"path": "/a.bin", "include_binary": True})

    assert "larger than max_bytes" in text and "as_link=true" in text
    assert response.asked == [100_004]  # never the whole file


async def test_include_binary_still_returns_text_that_fits(call, fake_fulcra):
    # 80 KB of text fits in max_bytes even though its base64 (107 KB) wouldn't:
    # text is returned as text.
    body = "ü" * 40_000  # 80,000 bytes of UTF-8
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, name="notes.txt", size=80_000)]
    _download_returns(fake_fulcra, body.encode())

    text = await call("read_file", {"path": "/notes.txt", "include_binary": True, "max_bytes": 100_000})

    assert text == f"Content of /notes.txt: {body}"


async def test_read_file_on_other_binaries_says_what_to_do(call, fake_fulcra):
    fake_fulcra.resolve_filepath.return_value = [dict(FILE, name="photo.jpg")]
    _download_returns(fake_fulcra, b"\xff\xd8\xff\xe0")

    text = await call("read_file", {"path": "/photo.jpg"})

    assert "as_link=true" in text
    assert "fulcra://files/photo.jpg" in text
    assert "include_binary" in text


async def test_write_file_takes_base64_and_guesses_the_type(call, fake_fulcra):
    fake_fulcra.upload_file.return_value = {"file": FILE}

    await call(
        "write_file",
        {"path": "/reports/q3.xlsx", "content": base64.b64encode(b"\x00\xffxl").decode(), "encoding": "base64"},
    )

    fake_fulcra.upload_file.assert_called_once_with(
        b"\x00\xffxl",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        4,
        "/reports/q3.xlsx",
    )


async def test_write_file_guesses_text_types(call, fake_fulcra):
    fake_fulcra.upload_file.return_value = {"file": FILE}

    await call("write_file", {"path": "/state.json", "content": "{}"})

    fake_fulcra.upload_file.assert_called_once_with(b"{}", "application/json", 2, "/state.json")


async def test_write_file_refuses_bad_or_large_base64(call, fake_fulcra, monkeypatch):
    bad = await call("write_file", {"path": "/a.bin", "content": "not base64!", "encoding": "base64"})
    monkeypatch.setattr(ft, "MAX_WRITE_BASE64_BYTES", 2)
    large = await call("write_file", {"path": "/a.bin", "content": "AAAA", "encoding": "base64"})

    assert "isn't valid base64" in bad
    assert "start_file_upload" in large and "open_file_manager" in large
    fake_fulcra.upload_file.assert_not_called()


# --- upload links (write_file encoding="upload_link") --------------------------------


async def test_write_file_hands_out_an_upload_link(call, fake_fulcra, hosted):
    text = await call("write_file", {"path": "/reports/q3.xlsx", "encoding": "upload_link"})

    assert "curl -sS -T <local file> 'https://mcp.example/files/upload/" in text
    # what to do if the sandbox blocks the server: names the link's own host
    assert "the sandbox is blocking mcp.example." in text
    assert "Settings > Capabilities > Domain allowlist" in text
    token = text.split("/files/upload/")[1].split("'")[0]
    record = ft._load_record("upload_links", token)
    assert record["path"] == "/reports/q3.xlsx" and record["grant_id"] == "grant-1"
    assert token not in "".join(p.name for p in settings.state_path.rglob("*"))
    fake_fulcra.upload_file.assert_not_called()


async def test_upload_link_mode_refuses_content_and_local_servers(call, fake_fulcra, hosted, monkeypatch):
    with_content = await call(
        "write_file", {"path": "/a.bin", "content": "x", "encoding": "upload_link"}
    )
    monkeypatch.setattr(settings, "fulcra_environment", "stdio")
    local = await call("write_file", {"path": "/a.bin", "encoding": "upload_link"})

    assert "leave content empty" in with_content
    assert "uvx fulcra-api file upload" in local
    fake_fulcra.upload_file.assert_not_called()


@pytest.fixture
def link_uploads(monkeypatch, put_chunks):
    """An upload link, with Fulcra's side faked: sessions and chunk storage."""
    import fulcra_mcp.main as main

    sessions = []

    def fake_session(fulcra, path, size, content_type):
        sessions.append((path, size, content_type))
        return {"url": "https://storage.example/up", "file": dict(FILE, size=size)}

    monkeypatch.setattr(ft, "_create_upload_session", fake_session)
    monkeypatch.setattr(main.oauth_provider, "reload_grant", lambda grant_id: object())
    monkeypatch.setattr(main, "fulcra_for_grant", lambda grant_id, creds: object())
    monkeypatch.setattr(settings, "oidc_server_url", "https://mcp.example")
    token = ft.make_upload_link("/reports/q3.xlsx", "grant-1").rsplit("/", 1)[1]
    return TestClient(main.app), token, sessions


def test_a_file_sent_to_an_upload_link_is_saved_in_chunks(link_uploads, put_chunks):
    client, token, sessions = link_uploads
    data = bytes(range(256)) * (ft.LINK_UPLOAD_CHUNK_BYTES // 256 + 5) + b"tail"

    response = client.put(f"/files/upload/{token}", content=data)

    assert response.status_code == 200, response.text
    assert response.json()["path"] == "/reports/q3.xlsx"
    assert response.json()["size"] == len(data)
    assert sessions == [(
        "/reports/q3.xlsx",
        len(data),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )]
    assert b"".join(chunk for _, chunk, _, _ in put_chunks) == data
    # every chunk but the last is a whole number of storage granules
    assert all(len(chunk) % ft.UPLOAD_GRANULE == 0 for _, chunk, _, _ in put_chunks[:-1])


def test_a_link_upload_resends_what_storage_did_not_confirm(link_uploads, monkeypatch):
    client, token, _ = link_uploads
    data = bytes(range(256)) * (ft.LINK_UPLOAD_CHUNK_BYTES // 256 + 5) + b"tail"
    stored = bytearray()
    calls = []

    def half_put(url, chunk, start, total):
        # Like storage under load: confirm half of each of the first few
        # requests, then everything.
        calls.append((start, len(chunk)))
        assert start == len(stored), "a request skipped bytes storage never confirmed"
        stored.extend(chunk if len(calls) > 3 else chunk[: max(len(chunk) // 2, 1)])
        return total if len(stored) == total else len(stored)

    monkeypatch.setattr(ft, "_put_chunk", half_put)
    response = client.put(f"/files/upload/{token}", content=data)

    assert response.status_code == 200, response.text
    assert bytes(stored) == data
    assert len(calls) > 3


def test_a_link_upload_that_storage_does_not_advance_fails_cleanly(link_uploads, monkeypatch):
    client, token, _ = link_uploads
    monkeypatch.setattr(ft, "_put_chunk", lambda url, chunk, start, total: start)

    response = client.put(f"/files/upload/{token}", content=b"abc")

    assert response.status_code == 502 and "Try again" in response.text
    assert ft._load_record("upload_links", token)["state"] == "new"


def test_resending_to_a_used_link_returns_the_same_result(link_uploads, put_chunks):
    client, token, sessions = link_uploads

    first = client.put(f"/files/upload/{token}", content=b"abc")
    again = client.put(f"/files/upload/{token}", content=b"abc")

    assert first.json() == again.json()
    assert len(sessions) == 1 and len(put_chunks) == 1


def test_an_empty_file_can_be_sent_to_a_link(link_uploads, put_chunks):
    client, token, sessions = link_uploads

    response = client.put(f"/files/upload/{token}", content=b"")

    assert response.status_code == 200 and response.json()["size"] == 0
    assert put_chunks[0][1:] == (b"", 0, 0)


def test_an_upload_link_needs_a_length_within_the_limit(link_uploads, put_chunks):
    client, token, _ = link_uploads

    def chunked():
        yield b"abc"

    no_length = client.put(f"/files/upload/{token}", content=chunked())
    too_big = client.put(
        f"/files/upload/{token}",
        content=b"x",
        headers={"Content-Length": str(ft.MAX_LINK_UPLOAD_BYTES + 1)},
    )

    assert no_length.status_code == 411 and "curl -T" in no_length.text
    assert too_big.status_code == 413 and "open_file_manager" in too_big.text
    assert put_chunks == []


@pytest.mark.parametrize(
    "length, pieces, message",
    [
        (10, [b"abc"], "ended before its Content-Length"),
        (2, [b"abc"], "more bytes than its Content-Length"),
    ],
)
async def test_a_body_that_disagrees_with_its_length_is_refused(
    link_uploads, put_chunks, length, pieces, message
):
    _, token, _ = link_uploads

    async def body():
        for piece in pieces:
            yield piece

    with pytest.raises(ft.LinkUploadError, match=message):
        await ft.receive_link_upload(token, length, body(), lambda grant_id: object())
    # the link can be used again
    assert ft._load_record("upload_links", token)["state"] == "new"


def test_a_failed_upload_can_be_retried_with_the_same_link(link_uploads, monkeypatch):
    client, token, sessions = link_uploads

    def broken(url, data, start, total):
        raise OSError("storage unreachable")

    monkeypatch.setattr(ft, "_put_chunk", broken)
    failed = client.put(f"/files/upload/{token}", content=b"abc")
    monkeypatch.setattr(ft, "_put_chunk", lambda url, data, start, total: start + len(data))
    retried = client.put(f"/files/upload/{token}", content=b"abc")

    assert failed.status_code == 502 and "Try again" in failed.text
    assert retried.status_code == 200


def test_an_upload_already_in_progress_is_refused(link_uploads, put_chunks):
    client, token, _ = link_uploads
    record = ft._load_record("upload_links", token)
    record.update(state="uploading", started_at=time.time())
    ft._save_record("upload_links", token, record)

    response = client.put(f"/files/upload/{token}", content=b"abc")

    assert response.status_code == 409 and put_chunks == []


def test_an_expired_or_ended_upload_link_says_so(link_uploads, monkeypatch, put_chunks):
    import fulcra_mcp.main as main

    client, token, _ = link_uploads
    unknown = client.put("/files/upload/not-a-token", content=b"abc")
    monkeypatch.setattr(main.oauth_provider, "reload_grant", lambda grant_id: None)
    ended = client.put(f"/files/upload/{token}", content=b"abc")

    assert unknown.status_code == 410 and "expired" in unknown.text
    assert ended.status_code == 410 and "Reconnect Fulcra" in ended.text
    assert put_chunks == []


# --- racing uploads to one link ----------------------------------------------------


class FakeBucket:
    """The state bucket seen through the GCS API, backed by the same files the
    tests' state directory holds (as gcsfuse would be). A file's generation is
    its mtime. ``after_read`` runs once a blob's contents have been read."""

    def __init__(self, after_read=lambda: None):
        import threading

        self.lock = threading.Lock()
        self.after_read = after_read

    def get_blob(self, name):
        path = settings.state_path / name
        return FakeBlob(self, path) if path.exists() else None


class FakeBlob:
    def __init__(self, bucket, path):
        self.bucket, self.path = bucket, path
        self.generation = path.stat().st_mtime_ns

    def _check(self, generation):
        from google.api_core.exceptions import PreconditionFailed

        if self.path.stat().st_mtime_ns != generation:
            raise PreconditionFailed("generation mismatch")

    def download_as_bytes(self, if_generation_match):
        with self.bucket.lock:
            self._check(if_generation_match)
            data = self.path.read_bytes()
        self.bucket.after_read()
        return data

    def upload_from_string(self, data, content_type, if_generation_match):
        with self.bucket.lock:
            self._check(if_generation_match)
            time.sleep(0.001)  # so the write moves the mtime on
            self.path.write_text(data)


def _race_two_uploads(token):
    """Send two uploads to one link at once; each one's outcome."""
    import asyncio
    import threading

    outcomes = [None, None]

    def send(i):
        async def body():
            yield b"abc"

        try:
            outcomes[i] = asyncio.run(
                ft.receive_link_upload(token, 3, body(), lambda grant_id: object())
            )
        except ft.LinkUploadError as e:
            outcomes[i] = e.status

    threads = [threading.Thread(target=send, args=(i,)) for i in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=10)
    return outcomes


def _assert_one_upload(outcomes, sessions, put_chunks):
    assert len(sessions) == 1 and len(put_chunks) == 1
    # the other request was refused, or (if it came after) got the same result
    results = [o for o in outcomes if isinstance(o, dict)]
    assert results and all(r == results[0] for r in results)
    assert all(o == 409 for o in outcomes if not isinstance(o, dict))


def test_racing_uploads_to_a_link_store_one_file(link_uploads, put_chunks, monkeypatch):
    import threading

    _, token, sessions = link_uploads
    # Both requests pass the first look at the link before either claims it.
    both_checked = threading.Barrier(2)
    change = ft._change_record

    def change_after_both_checked(*args):
        both_checked.wait(timeout=5)
        return change(*args)

    monkeypatch.setattr(ft, "_change_record", change_after_both_checked)

    _assert_one_upload(_race_two_uploads(token), sessions, put_chunks)


def test_racing_uploads_across_instances_store_one_file(link_uploads, put_chunks, monkeypatch):
    import threading

    _, token, sessions = link_uploads
    # Both instances read the same generation of the link before either writes.
    both_read = threading.Barrier(2)
    monkeypatch.setattr(settings, "state_bucket", "state")
    bucket = FakeBucket(lambda: both_read.wait(timeout=5))
    monkeypatch.setattr(ft, "_state_bucket", lambda: bucket)

    outcomes = _race_two_uploads(token)

    _assert_one_upload(outcomes, sessions, put_chunks)
    assert 409 in outcomes


def test_a_bucket_link_record_that_is_missing_or_expired_is_none(monkeypatch):
    monkeypatch.setattr(settings, "state_bucket", "state")
    bucket = FakeBucket()
    monkeypatch.setattr(ft, "_state_bucket", lambda: bucket)
    ft._save_record("upload_links", "old", {"state": "new", "expires_at": time.time() - 1})

    assert ft._change_record("upload_links", "missing", lambda r: r) is None
    assert ft._change_record("upload_links", "old", lambda r: r) is None
